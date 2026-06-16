import os
import json
import google.generativeai as genai
import docx
import openpyxl
from dotenv import load_dotenv

load_dotenv()

def extract_text_from_docx(file_path):
    doc = docx.Document(file_path)
    text = []
    # Tabloları oku
    for table in doc.tables:
        for row in table.rows:
            row_data = []
            for cell in row.cells:
                row_data.append(cell.text.strip().replace('\n', ' '))
            text.append(" | ".join(row_data))
    return "\n".join(text)

def extract_text_from_xlsx(file_path):
    wb = openpyxl.load_workbook(file_path, data_only=True)
    ws = wb.active
    text = []
    for row in ws.iter_rows(values_only=True):
        row_data = [str(cell).strip().replace('\n', ' ') if cell is not None else "" for cell in row]
        if any(row_data):
            text.append(" | ".join(row_data))
    return "\n".join(text)

def process_file_with_gemini(file_path):
    api_key = os.getenv("GEMINI_API_KEY")
    if not api_key:
        raise ValueError("GEMINI_API_KEY bulunamadı. Lütfen .env dosyasına ekleyin.")
        
    genai.configure(api_key=api_key)
        
    if file_path.lower().endswith('.docx'):
        raw_text = extract_text_from_docx(file_path)
    elif file_path.lower().endswith('.xlsx'):
        raw_text = extract_text_from_xlsx(file_path)
    else:
        raise ValueError("Desteklenmeyen dosya formatı.")

    prompt = f"""
Sen bir öğretmensin. Aşağıda bir öğretmenin eski yıllık planının ham tablosu verilmiştir.
Görev: Bu metindeki verileri okuyup, ders içeriğini satır satır sırayla ayıklayarak saf bir JSON dizisi oluşturmak.

Eğer Yöntem/Teknik, Materyal veya Açıklama boşsa boş string ("") bırak.
Hafta veya Tarih bilgisini göz ardı et (sadece içerik verilerini al).
Her bir veri satırı (haftalık içerik) için şu JSON formatını KESİNLİKLE koru:
[
  {{
    "kazanimlar": "...",
    "konular": "...",
    "yontem": "...",
    "materyaller": "...",
    "aciklama": "..."
  }}
]

Dikkat:
1. İçerik sadece kazanımlar ve konulardan ibarettir, okul adları veya alakasız verileri atla.
2. SADECE GEÇERLİ BİR JSON dondur, kod blogu (```json) kullanma, fazladan kelime etme.

Ham Tablo Verisi (Sınırlandırılmış):
{raw_text[:35000]}
"""

    model = genai.GenerativeModel('gemini-1.5-flash')
    response = model.generate_content(
        prompt,
        generation_config=genai.types.GenerationConfig(
            response_mime_type="application/json",
            temperature=0.1
        )
    )
    
    try:
        data = json.loads(response.text)
        return data
    except json.JSONDecodeError:
        text = response.text.replace('```json', '').replace('```', '').strip()
        return json.loads(text)
